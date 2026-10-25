from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from zipfile import ZipFile

source = Path(r'C:\Users\11c16\Downloads\medical_ai_capstone_submission.docx')
out = Path(r'C:\Users\11c16\OneDrive\Desktop\Deep-learning-capstome\medical_ai_capstone_submission_with_authors.docx')
doc = Document(source)
anchor = next((p for p in doc.paragraphs if p.text.strip() == 'Capstone Project Team'), None)
if anchor is None:
    raise RuntimeError('Could not find the submission author placeholder.')

people = [
    ('V Balasubramaniam', 'Computer Science (Data Science)', 'Vellore Institute of Technology', 'Chennai, Tamil Nadu', 'balasubramaniamv2005@gmail.com'),
    ('P Balaji Subramanian', 'Computer Science (Data Science)', 'Vellore Institute of Technology', 'Chennai, Tamil Nadu', 'spcamani@gmail.com'),
    ('Vijay V S', 'Computer Science (Data Science)', 'Vellore Institute of Technology', 'Nasiyanur, Tamil Nadu', 'vijay080504@gmail.com'),
]

def replace_paragraph(p, lines, is_new=False):
    p.clear()
    if not is_new:
        p.style = 'Author'
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(5)
    for i, line in enumerate(lines):
        if i:
            p.add_run().add_break()
        run = p.add_run(line)
        run.font.size = Pt(9)

# Use the template's five-line author block: name, department, institution,
# city/state, and email. Keep one clearly readable block per author.
replace_paragraph(anchor, people[0])
cursor = anchor._p
for person in people[1:]:
    p = doc.add_paragraph(style='Author')
    replace_paragraph(p, person, is_new=True)
    element = p._p
    element.getparent().remove(element)
    cursor.addnext(element)
    cursor = element

doc.save(out)
# Verify the saved file contains all author blocks before reporting success.
check = Document(out)
all_text = '\n'.join(p.text for p in check.paragraphs)
for name, _, _, _, email in people:
    if name not in all_text or email not in all_text:
        raise RuntimeError(f'Missing author data after save: {name}')
with ZipFile(out) as zf:
    if zf.testzip() is not None:
        raise RuntimeError('The saved DOCX package failed ZIP integrity validation.')
print(f'Saved: {out}')
print('Verified author paragraphs:')
for p in check.paragraphs[:5]:
    print(repr(p.text))
