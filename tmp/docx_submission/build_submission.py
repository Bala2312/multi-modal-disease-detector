from pathlib import Path
import copy, zipfile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT=Path.cwd()
TMP=ROOT/'tmp'/'docx_submission'
ORIGINAL=Path(r'C:\Users\11c16\Downloads\conference-template-a4.docx')
WORKING=TMP/'template_base.docx'
OUTPUT=ROOT/'medical_ai_capstone_submission.docx'
FIG=TMP/'architecture.png'

# Normalize the strict OOXML identifiers only in the task-local template copy.
nsmap={
 b'http://purl.oclc.org/ooxml/wordprocessingml/main':b'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
 b'http://purl.oclc.org/ooxml/officeDocument/relationships':b'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
 b'http://purl.oclc.org/ooxml/officeDocument/math':b'http://schemas.openxmlformats.org/officeDocument/2006/math',
 b'http://purl.oclc.org/ooxml/drawingml/wordprocessingDrawing':b'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
 b'http://purl.oclc.org/ooxml/drawingml/main':b'http://schemas.openxmlformats.org/drawingml/2006/main',
 b'http://purl.oclc.org/ooxml/drawingml/picture':b'http://schemas.openxmlformats.org/drawingml/2006/picture',
}
with zipfile.ZipFile(ORIGINAL) as src, zipfile.ZipFile(WORKING,'w',zipfile.ZIP_DEFLATED) as dst:
 for info in src.infolist():
  data=src.read(info.filename)
  for old,new in nsmap.items(): data=data.replace(old,new)
  # Transitional left/right alignment spellings are more portable than strict 'start'.
  if info.filename.endswith('.xml') or info.filename.endswith('.rels'):
   data=data.replace(b'w:val="start"',b'w:val="left"').replace(b'w:val="end"',b'w:val="right"')
  dst.writestr(info,data)

# Render a grayscale architecture figure in an IEEE-friendly style.
W,H=1100,960
im=Image.new('RGB',(W,H),'white'); draw=ImageDraw.Draw(im)
font_path=r'C:\Windows\Fonts\arial.ttf'; bold_path=r'C:\Windows\Fonts\arialbd.ttf'
def font(sz,bold=False):
 try: return ImageFont.truetype(bold_path if bold else font_path,sz)
 except OSError: return ImageFont.load_default()
def box(x1,y1,x2,y2,label,sub=None):
 draw.rounded_rectangle((x1,y1,x2,y2),radius=18,fill='#f3f6f8',outline='#23374d',width=3)
 lines=label.split('\n'); y=y1+14
 for line in lines:
  bb=draw.textbbox((0,0),line,font=font(27,True)); draw.text(((x1+x2-(bb[2]-bb[0]))/2,y),line,font=font(27,True),fill='#172338'); y+=36
 if sub:
  for line in sub.split('\n'):
   bb=draw.textbbox((0,0),line,font=font(20)); draw.text(((x1+x2-(bb[2]-bb[0]))/2,y+2),line,font=font(20),fill='#40536a'); y+=27
def arrow(x1,y1,x2,y2):
 draw.line((x1,y1,x2,y2),fill='#31485f',width=4)
 import math
 ang=math.atan2(y2-y1,x2-x1); L=15
 p1=(x2-L*math.cos(ang-math.pi/6),y2-L*math.sin(ang-math.pi/6)); p2=(x2-L*math.cos(ang+math.pi/6),y2-L*math.sin(ang+math.pi/6))
 draw.polygon([(x2,y2),p1,p2],fill='#31485f')
box(270,28,830,112,'Medical image upload','PNG / JPEG · validate dimensions and size')
box(350,150,750,230,'Modality routing','explicit override or image-statistics heuristic')
arrow(550,112,550,150)
box(45,285,510,395,'Chest X-ray model','DenseNet-121 · 14 NIH disease logits')
box(590,285,1055,395,'Skin lesion model','EfficientNet-B4 · current scaffold: 7 outputs')
arrow(440,230,285,285); arrow(660,230,815,285)
box(300,445,800,525,'Independent sigmoid probabilities','select top-scoring label for downstream query')
arrow(285,395,440,445); arrow(815,395,660,445)
box(20,605,350,720,'Grad-CAM','target class map and image overlay')
box(385,605,715,720,'Guideline RAG','PDF chunks · BioBERT · Chroma')
box(750,605,1080,720,'Triage rules','fixed label thresholds')
arrow(450,525,185,605); arrow(550,525,550,605); arrow(650,525,915,605)
box(240,815,860,925,'FastAPI response','prediction · score map · triage · heatmap path · grounded summary')
arrow(185,720,395,815); arrow(550,720,550,815); arrow(915,720,705,815)
im.save(FIG,dpi=(300,300))

doc=Document(str(WORKING))
body=doc._element.body
# Remove all sample/template content while retaining the final section properties.
for child in list(body):
 if child.tag != qn('w:sectPr'):
  body.remove(child)
sect=doc.sections[-1]
sect.page_width=Inches(8.27); sect.page_height=Inches(11.69)
sect.left_margin=Pt(44.65); sect.right_margin=Pt(44.65); sect.top_margin=Pt(54); sect.bottom_margin=Pt(72)
sect.header_distance=Pt(36); sect.footer_distance=Pt(36)
cols=sect._sectPr.find(qn('w:cols'))
if cols is None: cols=OxmlElement('w:cols'); sect._sectPr.append(cols)
cols.set(qn('w:num'),'2'); cols.set(qn('w:space'),'360')
# Use an IEEE-like font hierarchy from the retained template.
for stylename in ['Normal','Body Text','paper title','Author','Abstract','Keywords','Heading 1','Heading 2','Heading 3','Heading 5','references','figure caption','table head']:
 st=doc.styles[stylename]
 st.font.name='Times New Roman'
 st.font.color.rgb=RGBColor(0,0,0)
 st._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman')
 st._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman')
 st._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
# Keep heading hierarchy and captions together.
for stylename in ['Heading 1','Heading 2','Heading 3','Heading 5','table head','figure caption']:
 doc.styles[stylename].paragraph_format.keep_with_next=True

# Remove IEEE copyright placeholder from the retained footer; no fake publication metadata.
for s in doc.sections:
 for part in [s.footer,s.first_page_footer]:
  for p in part.paragraphs:
   for run in list(p.runs): run.text=''


def para(text='',style='Body Text',align=None,boldlead=None):
 p=doc.add_paragraph(style=style)
 if align is not None: p.alignment=align
 if boldlead and text.startswith(boldlead):
  p.add_run(boldlead).bold=True; p.add_run(text[len(boldlead):])
 else: p.add_run(text)
 p.paragraph_format.widow_control=True
 return p

def heading(text,level=1):
 p=doc.add_paragraph(style='Heading 1' if level==1 else 'Heading 2')
 p.add_run(text)
 p.paragraph_format.keep_with_next=True
 return p

def set_cell_shading(cell,fill):
 tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def set_cell_borders(cell,color='B7C0C8',sz='5'):
 tcPr=cell._tc.get_or_add_tcPr(); borders=tcPr.first_child_found_in('w:tcBorders')
 if borders is None: borders=OxmlElement('w:tcBorders'); tcPr.append(borders)
 for edge in ('top','left','bottom','right','insideH','insideV'):
  el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),sz); el.set(qn('w:color'),color); borders.append(el)
def add_table(headers,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(headers)); t.autofit=False
 try: t.style='Table Grid'
 except KeyError: pass
 for j,label in enumerate(headers):
  c=t.rows[0].cells[j]; c.text=label; set_cell_shading(c,'D9E2EA'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  for r in c.paragraphs[0].runs: r.bold=True; r.font.name='Times New Roman'; r.font.size=Pt(7.4)
  set_cell_borders(c)
 for i,row in enumerate(rows):
  cells=t.add_row().cells
  for j,value in enumerate(row):
   cells[j].text=value; cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   if i%2: set_cell_shading(cells[j],'F3F6F8')
   set_cell_borders(cells[j])
   for p in cells[j].paragraphs:
    p.paragraph_format.space_after=Pt(1); p.paragraph_format.space_before=Pt(1)
    for r in p.runs: r.font.name='Times New Roman'; r.font.size=Pt(7.2)
 if widths:
  for row in t.rows:
   for j,w in enumerate(widths): row.cells[j].width=Inches(w)
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(1); p.paragraph_format.space_before=Pt(0)
 return t

# Title block remains full width, as in the conference template.
p=doc.add_paragraph(style='paper title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Multi-Modal Medical Disease Detector and AI Clinical Assistant')
p=doc.add_paragraph(style='Author'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Capstone Project Team')

abstract=('This paper presents a modular research prototype for image-based clinical decision support across chest radiographs and dermoscopic images. A modality router dispatches images to a DenseNet-121 chest X-ray classifier or an EfficientNet-B4 skin-lesion classifier. Each network produces independent sigmoid probabilities. The pipeline adds class-specific Grad-CAM overlays, local retrieval over supplied medical guideline PDFs using BioBERT embeddings and ChromaDB, deterministic prototype triage rules, and a language-model summary constrained to retrieved passages. A FastAPI endpoint coordinates the workflow and returns the prediction, score map, triage output, heatmap path, and summary. The training scaffold uses focal loss, AdamW, cosine warm restarts, automatic mixed precision, patient-grouped multilabel folds for NIH data, and stratified folds for ISIC. The project pitch specifies targets of 0.92 NIH ROC-AUC, 0.88 ISIC sensitivity, 0.81 Grad-CAM localization IoU, and 0.96 retrieval factuality; these are goals rather than measured results. The current ISIC label mapping omits two challenge categories and maps actinic keratosis to a seven-class API vocabulary, so it must be reconciled before a full benchmark. Clinical validation, threshold calibration, and external testing remain outstanding.')
p=doc.add_paragraph(style='Abstract')
r=p.add_run('Abstract—'); r.bold=True; r.italic=True
p.add_run(abstract)
p=doc.add_paragraph(style='Keywords'); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
r=p.add_run('Index Terms—'); r.bold=True
p.add_run('chest radiography, dermoscopy, multi-label classification, Grad-CAM, retrieval-augmented generation, clinical decision support.')
# Transition from the full-width front matter to the two-column technical body.
breakp=doc.add_paragraph()
breakp.paragraph_format.space_after=Pt(0); breakp.paragraph_format.space_before=Pt(0)
pPr=breakp._p.get_or_add_pPr(); first=copy.deepcopy(sect._sectPr)
first.set(qn('w:rsidR'),'00000000')
sect_type=first.find(qn('w:type'))
if sect_type is None: sect_type=OxmlElement('w:type'); first.insert(0,sect_type)
sect_type.set(qn('w:val'),'continuous')
first_cols=first.find(qn('w:cols'))
if first_cols is None: first_cols=OxmlElement('w:cols'); first.append(first_cols)
first_cols.set(qn('w:num'),'1'); first_cols.set(qn('w:space'),'360')
pPr.append(first)

heading('I. INTRODUCTION',1)
para('Medical imaging systems can return classifications quickly, but a probability score alone does not show what image regions influenced the output or which supplied clinical references support a plain-language explanation. This project proposes a software pipeline that connects image classification, visual attribution, guideline retrieval, and deterministic triage in one API. The intended user is a clinician or research reviewer who can inspect the outputs; the system is not intended to diagnose patients autonomously.')
para('The term multi-modal in this work means that the service accepts two image modalities and routes them to separate models. It does not mean that image pixels and text are fused inside one neural network. Chest X-rays and dermoscopy have different acquisition characteristics and label spaces, so the prototype maintains separate classifiers and combines their downstream outputs only at the orchestration layer.')
para('The contributions are: (1) a routed dual-backbone classifier for 14 NIH chest findings and a restricted skin-label vocabulary; (2) class-specific Grad-CAM maps; (3) local guideline retrieval and a context-only summary prompt; and (4) deterministic triage rules exposed through a FastAPI endpoint. The current contribution is an architecture and software scaffold. The supplied proposal does not provide completed clinical experiments, and no performance results are claimed here.')

heading('II. RELATED WORK',1)
heading('A. Medical image classification',2)
para('The NIH ChestX-ray dataset popularized weakly supervised, multi-label recognition of thoracic findings from frontal radiographs [1]. DenseNet introduced direct connections among convolutional layers to encourage feature reuse and improve information flow [2]. EfficientNet uses compound scaling across network depth, width, and input resolution [3]. These families provide a practical transfer-learning baseline for the two image domains; their selection does not establish suitability for a clinical population without local validation.')
para('The ISIC 2019 challenge defines eight known lesion diagnoses in its training labels, with an additional unknown/outlier category in the test setting [4]. This project currently maps the source AK column to an API label AKIEC, omits SCC, and excludes UNK. This changes the target definition and prevents claiming full nine-category challenge performance until the mapping is revised and the model head is retrained.')
heading('B. Imbalance and explanation',2)
para('Focal loss reshapes cross-entropy so confidently classified examples contribute less loss, emphasizing harder examples under class imbalance [5]. In this project it is applied elementwise to independent label logits. Grad-CAM uses gradients with respect to a late convolutional feature tensor to create a coarse, class-specific localization map [6]. It can expose image regions associated with a score, but it is not a segmentation model and does not guarantee that the highlighted tissue is clinically causal.')
heading('C. Evidence-grounded language generation',2)
para('Retrieval-augmented generation conditions text generation on retrieved non-parametric sources [7]. The prototype indexes local PDF guideline chunks with BioBERT embeddings [8] and queries ChromaDB. Retrieval and a strict prompt can narrow the available evidence; they do not guarantee factuality. Source quality, chunking, embedding suitability, retrieval errors, and model behavior must all be evaluated.')

heading('III. DATA AND MODEL METHODS',1)
heading('A. Dataset and target definitions',2)
para('The chest dataset is NIH ChestX-ray14, with a 14-element target vector derived from the Finding Labels field. Multiple findings can be positive for one image. “No Finding” is not a model output in the current API-compatible head; such images receive all-zero disease targets. The skin dataset is ISIC 2019, whose official task includes eight known diagnostic categories plus an unknown/outlier category [4]. The present training code uses seven outputs (MEL, NV, BCC, AKIEC, BKL, DF, VASC), maps AK to AKIEC, and omits SCC and UNK. This label mismatch is an implementation limitation that must be fixed before a faithful ISIC 2019 evaluation.')
add_table(['Modality','Current output labels','Target caveat'],[
 ('Chest X-ray','14 NIH findings','No Finding is represented by an all-zero disease vector, not an output class.'),
 ('Skin lesion','MEL, NV, BCC, AKIEC, BKL, DF, VASC','AK is mapped to AKIEC; SCC and UNK are excluded. Not the complete ISIC 2019 label space.'),
],widths=[0.85,1.15,1.30])
heading('B. Preprocessing and augmentation',2)
para('Both data loaders decode images as RGB, resize to 224 × 224, convert them to tensors, and normalize channels with ImageNet mean (0.485, 0.456, 0.406) and standard deviation (0.229, 0.224, 0.225). Training augmentation includes random resized crop, horizontal flip, and random affine transformations. ColorJitter is applied to skin images; it is omitted for radiographs. The resize policy and augmentation range are initial engineering choices, not clinically validated invariances. In particular, horizontal flipping can alter laterality cues.')
heading('C. Backbones and multi-label head',2)
para('The chest branch uses torchvision DenseNet-121 and replaces its classifier with a 14-unit linear layer. The skin branch uses EfficientNet-B4 and replaces its final classifier with a seven-unit layer in the current scaffold. The models return logits z; inference applies the sigmoid independently, p_k = 1/(1 + exp(−z_k)). Probabilities are not normalized to sum to one, which permits co-occurring chest findings. ImageNet initialization is optional at training time. The shared 224-pixel input is a prototype setting; input resolution should be studied for EfficientNet-B4 and validated against the source image protocols.')
heading('D. Objective and optimization',2)
para('For a binary target y and logit z, let p = sigmoid(z), p_t = y p + (1−y)(1−p), and alpha_t = alpha y + (1−alpha)(1−y). The elementwise focal objective used by the training scaffold is:')
p=para('FL(p_t) = alpha_t (1 − p_t)^gamma BCEWithLogits(z, y).',style='Body Text'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
para('The factor (1−p_t)^gamma suppresses easy examples; alpha_t balances positive and negative terms. This loss can improve focus on hard examples but does not by itself correct dataset bias or guarantee minority-class sensitivity. Training uses AdamW with weight decay and CosineAnnealingWarmRestarts. Within a cycle, the learning rate follows eta_t = eta_min + 0.5(eta_max−eta_min)[1+cos(pi T_cur/T_i)]; the schedule returns to its cycle maximum at a restart and later cycles expand according to T_mult. CUDA autocast and GradScaler are enabled when a CUDA device is available.')

heading('IV. EXPLAINABILITY, RETRIEVAL, AND TRIAGE',1)
heading('A. Grad-CAM overlays',2)
para('For target class c and feature channel k, the spatially averaged gradient defines alpha_k^c = (1/Z) sum_i sum_j (partial y^c / partial A_ij^k). The map is L^c = ReLU(sum_k alpha_k^c A^k). The implementation captures the target-layer activation in a forward hook and registers a tensor gradient hook on the activation output. It uses features.denseblock4 for DenseNet-121 and features.7 for EfficientNet-B4, upsamples the normalized map to the original image size, applies OpenCV COLORMAP_JET, blends it with the source image, and writes the overlay to outputs/heatmaps/. The API returns a local path; it does not stream the heatmap bytes.')
heading('B. Grounded retrieval and summary',2)
para('PDF guideline pages placed in data/guidelines/ are split with RecursiveCharacterTextSplitter at 500 characters with 50-character overlap. HuggingFaceEmbeddings uses dmis-lab/biobert-base-cased-v1.1; ChromaDB stores a persistent local vector collection. The top predicted label is the retrieval query, and the top-K passages are included with source metadata in the LLM prompt. The prompt directs the model to use only those passages, treat them as untrusted text, identify insufficient context, avoid individualized treatment directions, and describe model outputs as unconfirmed. If no passages are retrieved, the API returns a deterministic insufficient-context message without constructing the LLM client. BioBERT is a biomedical language representation model rather than a sentence-embedding-trained retrieval model; retrieval quality must be measured on the actual guideline corpus.')
heading('C. Deterministic triage',2)
para('Triage is calculated from the probability map before text synthesis. The current rule table assigns emergency level to named triggers such as pneumothorax and melanoma at 0.50, standard consult to pneumonia at 0.35, and unlisted findings at or above a general 0.40 threshold. When several triggers are met, the most urgent level takes precedence. With no positive trigger the rule returns low risk; because the current NIH output vocabulary omits No Finding, that result does not represent a calibrated normal prediction. These thresholds are prototype software rules and must not be used to direct patient care.')

heading('V. API AND END-TO-END WORKFLOW',1)
para('The FastAPI POST /analyze endpoint accepts a multipart image upload and an optional modality override (chest_xray or skin_lesion). It checks upload size and pixel count, decodes supported images, and invokes lazy-loaded services. If no override is supplied, a simple image-saturation heuristic selects the branch. This heuristic can misroute grayscale dermoscopy, colorized radiographs, screenshots, and unusual image pipelines; a known modality should be passed explicitly.')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(FIG),width=Inches(3.25))
p=doc.add_paragraph(style='figure caption'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Fig. 1. Routed image analysis and downstream evidence workflow.')
para('The vision service returns the probability associated with each configured label. The orchestration layer selects the highest-scoring label for Grad-CAM and RAG query, computes triage from all label probabilities, and assembles the structured response: top diagnosis label, confidence, triage level and detail, heatmap path, clinical summary, modality, and full probability dictionary. The response is a research artifact and must be interpreted by a qualified reviewer.')

heading('VI. EXPERIMENTAL VALIDATION PLAN',1)
para('The project pitch states four evaluation goals: NIH macro ROC-AUC of 0.92, ISIC sensitivity of 0.88, Grad-CAM localization IoU of 0.81, and RAG factuality of 0.96 [9]. They are proposal targets only; the pitch provides no completed validation run, sample counts, confidence intervals, operating thresholds, or expert annotation protocol. They are therefore listed as objectives rather than results.')
add_table(['Evaluation target','Proposal goal','Required measurement'],[
 ('NIH chest classifier','ROC-AUC = 0.92','Per-label AUC on patient-separated test data; macro aggregation rules stated.'),
 ('ISIC classifier','Sensitivity = 0.88','Per-class sensitivity, macro summary, confusion matrix, and calibrated thresholds.'),
 ('Grad-CAM localization','IoU = 0.81','Compare thresholded heatmaps with expert masks; specify mask availability and protocol.'),
 ('RAG factuality','0.96','Clinician-reviewed claim support, source coverage, and unsupported-claim rate.'),
],widths=[0.95,0.75,1.60])
para('The current training scaffold performs K-fold validation, uses patient-grouped multilabel stratification for NIH, and stratifies ISIC folds by class. It early-stops each fold on validation macro ROC-AUC and records the best fold checkpoint. Selecting the best fold from the same cross-validation scores is not an unbiased final estimate. A submission-grade evaluation should lock a patient/lesion-separated external test set before model selection and report confidence intervals, class-wise metrics, calibration, and threshold selection on validation data only.')
para('Ablations should compare a backbone with deterministic preprocessing against (1) augmentation, (2) focal loss versus BCEWithLogits, and (3) the selected training schedule. Grad-CAM localization and RAG grounding require separate expert-reviewed protocols. Resource evaluation should report hardware, batch size, model load time, end-to-end latency, peak memory, and CPU/GPU utilization. No ablation or resource measurements are currently available.')

heading('VII. LIMITATIONS AND SAFETY',1)
para('The data labels are weak labels and may contain reporting noise. Public benchmark performance may not transfer across institutions, equipment, prevalence, or patient populations. The current skin-label reduction is not equivalent to the complete ISIC 2019 task; the X-ray head has no explicit No Finding class. The modality heuristic is a coarse routing cue. A 224-pixel resize may discard small details. Focal loss and cross-validation do not resolve these dataset limitations.')
para('Grad-CAM is a coarse attribution map and may highlight confounders. Retrieved PDF text can be incomplete or outdated, and a prompt cannot guarantee that every generated claim is supported. Triage cutoffs have not been clinically calibrated. Before any prospective use, the team must reconcile labels, obtain a held-out external cohort, establish a clinician-reviewed reference standard, test subgroup performance, calibrate outputs, validate retrieval and generated claims, and implement security, privacy, monitoring, and governance controls. The prototype must not be used for autonomous diagnosis, treatment, or emergency decisions.')

heading('VIII. CONCLUSION',1)
para('This work specifies a modular medical image decision-support pipeline that routes chest radiographs and dermoscopic images to separate convolutional classifiers, adds class-specific Grad-CAM, retrieves local guideline passages, and applies deterministic prototype triage before producing a structured API response. Its design makes each stage inspectable and supports a reproducible validation plan. The evidence currently supports an architecture proposal and implementation scaffold, not a claim of clinical accuracy. The immediate research priorities are to correct label definitions, train and evaluate on leakage-controlled splits, and validate every explanation, retrieval, and triage component with independent clinical review.')

heading('REFERENCES',1)
references=[
'[1] X. Wang et al., “ChestX-ray8: Hospital-scale chest X-ray database and benchmarks on weakly-supervised classification and localization of common thorax diseases,” in Proc. IEEE Conf. Computer Vision and Pattern Recognition, 2017, doi: 10.1109/CVPR.2017.369.',
'[2] G. Huang, Z. Liu, L. van der Maaten, and K. Q. Weinberger, “Densely connected convolutional networks,” in Proc. IEEE Conf. Computer Vision and Pattern Recognition, 2017, pp. 4700–4708. doi: 10.1109/CVPR.2017.243.',
'[3] M. Tan and Q. V. Le, “EfficientNet: Rethinking model scaling for convolutional neural networks,” in Proc. 36th Int. Conf. Machine Learning, PMLR, vol. 97, 2019, pp. 6105–6114.',
'[4] International Skin Imaging Collaboration, “ISIC 2019 Challenge and data,” ISIC Archive. [Online]. Available: https://challenge.isic-archive.com/landing/2019/ (accessed Oct. 23, 2026).',
'[5] T.-Y. Lin, P. Goyal, R. Girshick, K. He, and P. Dollár, “Focal loss for dense object detection,” in Proc. IEEE Int. Conf. Computer Vision, 2017, doi: 10.1109/ICCV.2017.324.',
'[6] R. R. Selvaraju et al., “Grad-CAM: Visual explanations from deep networks via gradient-based localization,” in Proc. IEEE Int. Conf. Computer Vision, 2017, pp. 618–626, doi: 10.1109/ICCV.2017.74.',
'[7] P. Lewis et al., “Retrieval-augmented generation for knowledge-intensive NLP tasks,” in Advances in Neural Information Processing Systems, vol. 33, 2020.',
'[8] J. Lee et al., “BioBERT: A pre-trained biomedical language representation model for biomedical text mining,” Bioinformatics, vol. 36, no. 4, pp. 1234–1240, 2020, doi: 10.1093/bioinformatics/btz682.',
'[9] Capstone Project Team, “Multi-Modal Medical Disease Detector and AI Clinical Assistant: Project Pitch,” project presentation, 2026.'
]
for ref in references:
 p=para(ref,style='references'); p.paragraph_format.left_indent=Inches(.18); p.paragraph_format.first_line_indent=Inches(-.18)

# Refresh title/author metadata while leaving the retained template structure intact.
doc.core_properties.title='Multi-Modal Medical Disease Detector and AI Clinical Assistant'
doc.core_properties.subject='Conference-style capstone project paper'
doc.core_properties.author='Capstone Project Team'
doc.core_properties.keywords='medical imaging, deep learning, Grad-CAM, retrieval-augmented generation'
doc.save(OUTPUT)
print('Wrote',OUTPUT,'bytes',OUTPUT.stat().st_size)
print('Paragraphs',len(doc.paragraphs),'tables',len(doc.tables),'sections',len(doc.sections))
