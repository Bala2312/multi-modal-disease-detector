from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

source = Path(r'C:\Users\11c16\Downloads\medical_ai_capstone_submission.docx')
out = Path(r'C:\Users\11c16\OneDrive\Desktop\Deep-learning-capstome\medical_ai_capstone_submission.docx')
doc = Document(source)
anchor = next((p for p in doc.paragraphs if p.text.strip() == 'Capstone Project Team'), None)
if anchor is None:
    raise RuntimeError('Expected author placeholder was not found; refusing to edit an unexpected document.')

section = doc.sections[0]
usable_width = section.page_width - section.left_margin - section.right_margin
col_width = int(usable_width / 3)

people = [
    ('V Balasubramaniam', 'Computer Science (Data Science)', 'Vellore Institute of Technology', 'Chennai, Tamil Nadu', 'balasubramaniamv2005@gmail.com'),
    ('P Balaji Subramanian', 'Computer Science (Data Science)', 'Vellore Institute of Technology', 'Chennai, Tamil Nadu', 'spcamani@gmail.com'),
    ('Vijay V S', 'Computer Science (Data Science)', 'Vellore Institute of Technology', 'Nasiyanur, Tamil Nadu', 'vijay080504@gmail.com'),
]

table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
for i, (cell, person) in enumerate(zip(table.rows[0].cells, people)):
    cell.width = col_width
    table.columns[i].width = col_width
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for j, line in enumerate(person):
        p = cell.paragraphs[0] if j == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(line)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(8.5)
        if j == 0:
            run.bold = True

# Keep the author row visually unobtrusive, as in conference author blocks.
tbl_pr = table._tbl.tblPr
borders = OxmlElement('w:tblBorders')
for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
    tag = OxmlElement(f'w:{edge}')
    tag.set(qn('w:val'), 'nil')
    borders.append(tag)
tbl_pr.append(borders)

# Place the author table exactly where the template's author placeholder sits.
table_el = table._tbl
body = doc._element.body
body.remove(table_el)
anchor._p.addnext(table_el)
anchor._p.getparent().remove(anchor._p)

doc.save(out)
print(f'Updated submission: {out}')
print(f'Author block width: {usable_width / 914400:.2f} in total, {col_width / 914400:.2f} in per author')

