from docx import Document
from pathlib import Path
paths=[r'C:\Users\11c16\OneDrive\Desktop\Deep-learning-capstome\tmp\docx_submission\template_base.docx',r'C:\Users\11c16\Downloads\medical_ai_capstone_submission.docx',r'C:\Users\11c16\OneDrive\Desktop\Deep-learning-capstome\medical_ai_capstone_submission.docx']
for path in paths:
 print('\nFILE',Path(path).name)
 d=Document(path); body=d._element.body
 for i,e in enumerate(list(body)[:18]):
  kind=e.tag.split('}')[-1]; text=''.join(e.itertext()).replace('\n',' ')[:450]; print(i,kind,repr(text))
 print('paragraphs',[(p.style.name,p.text[:200]) for p in d.paragraphs[:12]])
 print('tables',len(d.tables))
 for j,t in enumerate(d.tables[:2]): print('TABLE',j,len(t.rows),len(t.columns),[[c.text[:100] for c in r.cells] for r in t.rows[:3]])


